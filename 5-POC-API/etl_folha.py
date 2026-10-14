# -*- coding: utf-8 -*-
"""
Prova de conceito: motor de integracao Folha (origem) -> ERP (destino).

Objetivo desta PoC
------------------
Reproduzir, em codigo versionado e testavel, exatamente o que hoje a planilha
"Planilha para tratar e preparar dado para Importar Folha no RP Info.xlsx"
produz atraves de 8 consultas Power Query -- e, alem disso, aplicar as
validacoes que hoje sao feitas visualmente.

Criterio de aceite: a saida deve bater LINHA A LINHA com a aba "Importacao"
da planilha atual, na competencia 08/2026.

O que esta PoC NAO e
--------------------
Nao e a arquitetura final. As entradas aqui sao os mesmos arquivos que o
processo atual usa. Na solucao definitiva, `ler_folha_dominio` e
`ler_base_colaboradores` viram adaptadores de API/banco -- e so eles mudam.
"""

from __future__ import annotations

import csv
import sys
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import openpyxl

BASE = Path(__file__).resolve().parent.parent
SAIDA = Path(__file__).resolve().parent / "saida"

ARQ_FOLHA = BASE / "3-Atividade - Tratar Dado Folha para Importação" / \
    "Arquivo de Importação para Excel (Tratar Query Domínio Folha).TXT"
ARQ_PLANILHA = BASE / "3-Atividade - Tratar Dado Folha para Importação" / \
    "Planilha para tratar e preparar dado para Importar Folha no RP Info.xlsx"
ARQ_PLANOGER = BASE / "3-Atividade - Tratar Dado Folha para Importação" / "Plano Gerencial RP.txt"

# ---------------------------------------------------------------------------
# CADASTROS
#
# Hoje estes tres blocos estao cravados dentro do codigo M da planilha. Na
# solucao definitiva viram tabela com tela e historico de alteracao. Aqui ficam
# no topo do arquivo, isolados, justamente para mostrar o que precisa sair do
# codigo.
# ---------------------------------------------------------------------------

DE_PARA_LOJAS = {
    762: "001", 763: "002", 761: "003", 793: "004", 820: "005", 860: "006",
    899: "007", 924: "008", 1: "010", 2: "011", 3: "012", 1011: "013",
    1108: "014", 925: "031", 1159: "015", 1160: "016", 1181: "017", 1207: "018",
}

# Contas que pertencem a um terceiro (nao a um colaborador): o codigo do
# "funcionario" no destino e fixo, independente de quem gerou o lancamento.
CONTAS_DE_TERCEIRO = {
    211415: {"cnpj": "00394460005887", "descricao": "Receita Federal", "codigo_destino": None},
    211402: {"cnpj": "00360305000104", "descricao": "Caixa Economica Federal", "codigo_destino": 128384},
    331207: {"cnpj": "00360305000104", "descricao": "Caixa Economica Federal", "codigo_destino": 128384},
}

# Lancamentos de FGTS chegam da origem com debito E credito na MESMA linha.
# A leitura direta so aproveita a perna de debito; a perna de credito precisa
# ser recuperada. Ver `recuperar_pernas_de_credito`.
CONTA_CREDITO_DUPLA_FACE = 211402
CONTAS_DEBITO_A_ZERAR = {331207, 212401}

# ---------------------------------------------------------------------------
# PONTOS EM ABERTO
#
# Cada flag abaixo e uma pergunta que a documentacao NAO responde e que precisa
# ser confirmada com o cliente. O default reproduz o comportamento atual --
# porque o que esta em producao hoje e a unica referencia que temos.
# ---------------------------------------------------------------------------

# A origem entrega o historico com espaco a esquerda (" 08/2026"). O processo
# atual envia assim. Aparar e mais limpo, mas muda 234 linhas por competencia.
# PERGUNTA: o destino usa esse campo para conciliar algo?
APARAR_HISTORICO = False

# O indicador ordinal "º" NAO esta na lista de 24 letras que a planilha troca,
# entao ele chega ao destino como caractere nao-ASCII (61 linhas/mes).
# PERGUNTA: o destino aceita "º" ou isso vira lixo na tela dele?
NORMALIZAR_ORDINAL_MASCULINO = False

# ---------------------------------------------------------------------------
# MODELO
# ---------------------------------------------------------------------------


@dataclass
class Lancamento:
    loja_origem: int
    loja_destino: str | None
    data: date
    debito: int
    credito: int
    descricao_conta: str
    valor: Decimal
    historico: str
    nome_colaborador: str
    origem_linha: int
    perna_recuperada: bool = False
    codigo_colaborador: str | None = None
    codigo_final: str | None = None
    tipo_gerencial: str | None = None

    @property
    def conta(self) -> int:
        return self.credito if self.debito == 0 else self.debito

    @property
    def natureza(self) -> str:
        return "DB" if self.debito > 0 else "CR"


@dataclass
class Ocorrencia:
    severidade: str   # BLOQUEIO | ALERTA | INFO
    codigo: str
    mensagem: str
    detalhe: list = field(default_factory=list)


# ---------------------------------------------------------------------------
# ADAPTADOR DE ORIGEM
# ---------------------------------------------------------------------------


def _decimal_de_centavos(bruto: str) -> Decimal:
    """Origem entrega valor em centavos, sem separador, com zeros a esquerda."""
    return Decimal(bruto.strip()) / 100


def ler_folha_dominio(caminho: Path):
    """Adaptador de origem. Hoje le arquivo; amanha chama API. Mesmo retorno."""
    lancamentos = []
    ocorrencias = []
    descartadas = []

    with open(caminho, encoding="cp1252", newline="") as fh:
        for n, campos in enumerate(csv.reader(fh, delimiter=","), start=1):
            if len(campos) != 9:
                # Linha de cabecalho de bloco por loja. Nao e dado.
                continue
            try:
                loja = int(campos[1])
                data = datetime.strptime(campos[2].strip(), "%d/%m/%Y").date()
                debito = int(campos[3])
                credito = int(campos[4])
                valor = _decimal_de_centavos(campos[6])
            except (ValueError, ArithmeticError) as erro:
                # O processo atual apaga estas linhas em silencio.
                # Aqui elas viram BLOQUEIO.
                descartadas.append(f"linha {n}: {erro}")
                continue

            lancamentos.append(Lancamento(
                loja_origem=loja,
                loja_destino=DE_PARA_LOJAS.get(loja),
                data=data,
                debito=debito,
                credito=credito,
                descricao_conta=campos[5].strip(),
                valor=valor,
                historico=campos[7].strip() if APARAR_HISTORICO else campos[7].rstrip(),
                nome_colaborador=campos[8].strip().upper(),
                origem_linha=n,
            ))

    if descartadas:
        ocorrencias.append(Ocorrencia(
            "BLOQUEIO", "VALOR_ILEGIVEL",
            f"{len(descartadas)} linha(s) da origem nao puderam ser convertidas. "
            "O processo atual descartaria estas linhas sem aviso.",
            descartadas[:20],
        ))
    return lancamentos, ocorrencias


def recuperar_pernas_de_credito(lancamentos):
    """
    A origem emite o FGTS como uma unica linha com debito 331207 E credito
    211402. Lida de forma direta, a linha vira so o debito -- e o credito some,
    quebrando a partida dobrada.

    A planilha resolve isso relendo o arquivo inteiro do zero numa segunda
    consulta (~40 passos copiados). Aqui a mesma regra cabe em poucas linhas e
    ganha um nome.
    """
    recuperadas = []
    for lc in lancamentos:
        if lc.credito != CONTA_CREDITO_DUPLA_FACE or lc.debito not in CONTAS_DEBITO_A_ZERAR:
            continue
        gemea = Lancamento(**dict(lc.__dict__))
        gemea.debito = 0
        gemea.perna_recuperada = True
        recuperadas.append(gemea)
    return recuperadas


# ---------------------------------------------------------------------------
# CADASTROS AUXILIARES
# ---------------------------------------------------------------------------


def ler_base_colaboradores(caminho: Path):
    """
    Adaptador de identidade. Hoje: aba da planilha, alimentada por um dump
    completo (76 campos, ~5.600 pessoas) gravado em pasta de rede.
    Amanha: consulta por CPF na API do destino, sem dump.
    """
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    ws = wb["dFunc_RPMix"]
    linhas = ws.iter_rows(values_only=True)
    cab = list(next(linhas))
    col = {nome: i for i, nome in enumerate(cab) if nome}

    por_nome = defaultdict(list)
    total = 0
    sem_cpf = []
    for linha in linhas:
        if linha[col["Nome"]] is None:
            continue
        total += 1
        registro = {
            "codigo": str(linha[col["Código"]]).strip(),
            "nome": str(linha[col["Nome"]]).strip().upper(),
            "cpf": (str(linha[col["C.P.F."]]).strip() if linha[col["C.P.F."]] else ""),
            "pis": (str(linha[col["PIS"]]).strip() if linha[col["PIS"]] else ""),
        }
        if not registro["cpf"]:
            sem_cpf.append(registro["codigo"])
        por_nome[registro["nome"]].append(registro)
    wb.close()

    ocorrencias = []
    homonimos = {n: r for n, r in por_nome.items() if len(r) > 1}
    if homonimos:
        ocorrencias.append(Ocorrencia(
            "ALERTA", "HOMONIMOS_NA_BASE",
            f"{len(homonimos)} nome(s) aparecem em mais de um cadastro. "
            "Enquanto a conciliacao for por nome, cada um destes duplica o "
            "lancamento se o colaborador entrar na folha.",
            [f"{n} -> codigos {[x['codigo'] for x in r]}" for n, r in homonimos.items()],
        ))
    if total:
        ocorrencias.append(Ocorrencia(
            "INFO", "COBERTURA_CPF",
            f"{total - len(sem_cpf)} de {total} cadastros possuem CPF "
            f"({(total - len(sem_cpf)) / total:.1%}). A chave definitiva ja existe.",
        ))
    return dict(por_nome), ocorrencias


def ler_plano_gerencial(caminho: Path):
    plano = {}
    with open(caminho, encoding="cp1252", newline="") as fh:
        leitor = csv.reader(fh, delimiter=";")
        next(leitor)
        for linha in leitor:
            if len(linha) >= 2 and linha[0].strip():
                plano[int(linha[0])] = linha[1].strip()
    return plano


def ler_codigos_manuais(caminho: Path):
    """A aba 'Nomes Divergentes': onde a conciliacao por nome falha e alguem
    digita o codigo na mao. Na solucao definitiva, isto deve ficar vazio."""
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    ws = wb["Nomes Divergentes"]
    manuais = {}
    for linha in ws.iter_rows(min_row=2, values_only=True):
        if len(linha) >= 7 and linha[5] and linha[6]:
            manuais[str(linha[5]).strip().upper()] = str(linha[6]).strip()
    wb.close()
    return manuais


# ---------------------------------------------------------------------------
# TRANSFORMACAO
# ---------------------------------------------------------------------------


# As 24 letras que a planilha troca, uma a uma, no codigo M.
_TABELA_DA_PLANILHA = str.maketrans(
    "áàãâéêíóôõúçÁÀÃÂÉÊÍÓÔÕÚÇ",
    "aaaaeeioooucAAAAEEIOOOUC",
)


def sem_acento(texto: str) -> str:
    """
    O destino nao aceita acentuacao.

    A planilha troca 24 letras uma a uma -- e deixa passar tudo que nao esta
    na lista: 'º', 'ª', 'ü', 'ñ', 'Ü', 'Ñ'. A normalizacao Unicode (NFKD)
    cobre todas de uma vez.

    Enquanto a pergunta sobre o 'º' nao for respondida, o default reproduz o
    comportamento atual, para nao introduzir diferenca por conta propria.
    """
    if NORMALIZAR_ORDINAL_MASCULINO:
        return unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode("ascii")
    return texto.translate(_TABELA_DA_PLANILHA)


def residuo_nao_ascii(texto: str) -> set:
    """Caracteres que sobreviveram a normalizacao e vao chegar ao destino."""
    return {c for c in texto if ord(c) > 127}


def enriquecer(lancamentos, base_colaboradores, plano_gerencial, codigos_manuais):
    ocorrencias = []
    nao_resolvidos = Counter()
    conta_sem_plano = Counter()
    loja_sem_depara = Counter()
    por_nome = por_terceiro = por_digitacao = 0
    residuos = Counter()

    for lc in lancamentos:
        if lc.loja_destino is None:
            loja_sem_depara[lc.loja_origem] += 1

        lc.historico = sem_acento(lc.historico)
        for c in residuo_nao_ascii(lc.historico):
            residuos[c] += 1

        terceiro = CONTAS_DE_TERCEIRO.get(lc.conta)
        if terceiro and terceiro["codigo_destino"] is not None:
            lc.codigo_final = str(terceiro["codigo_destino"])
            por_terceiro += 1
        else:
            candidatos = base_colaboradores.get(lc.nome_colaborador, [])
            if len(candidatos) == 1:
                lc.codigo_colaborador = candidatos[0]["codigo"]
                lc.codigo_final = candidatos[0]["codigo"]
                por_nome += 1
            elif len(candidatos) > 1:
                ocorrencias.append(Ocorrencia(
                    "BLOQUEIO", "HOMONIMO_NA_FOLHA",
                    f"{lc.nome_colaborador} casa com {len(candidatos)} cadastros. "
                    "Impossivel decidir por nome.",
                    [c["codigo"] for c in candidatos],
                ))
            else:
                manual = codigos_manuais.get(lc.nome_colaborador)
                if manual:
                    lc.codigo_final = manual
                    por_digitacao += 1
                else:
                    nao_resolvidos[lc.nome_colaborador] += 1

        lc.tipo_gerencial = plano_gerencial.get(lc.conta)
        if lc.tipo_gerencial is None:
            conta_sem_plano[lc.conta] += 1

    if nao_resolvidos:
        ocorrencias.append(Ocorrencia(
            "BLOQUEIO", "COLABORADOR_NAO_RESOLVIDO",
            f"{len(nao_resolvidos)} colaborador(es) sem codigo no destino.",
            [f"{n} ({q} lancamentos)" for n, q in nao_resolvidos.most_common(20)],
        ))
    if conta_sem_plano:
        ocorrencias.append(Ocorrencia(
            "BLOQUEIO", "CONTA_SEM_PLANO_GERENCIAL",
            f"{len(conta_sem_plano)} conta(s) sem classificacao gerencial. "
            "Rubrica nova costuma aparecer aqui.",
            [f"conta {c} ({q} lancamentos)" for c, q in conta_sem_plano.most_common()],
        ))
    if loja_sem_depara:
        ocorrencias.append(Ocorrencia(
            "BLOQUEIO", "LOJA_SEM_DE_PARA",
            f"{len(loja_sem_depara)} loja(s) da origem sem equivalente no destino.",
            [f"loja origem {l} ({q} lancamentos)" for l, q in loja_sem_depara.most_common()],
        ))
    if residuos:
        ocorrencias.append(Ocorrencia(
            "ALERTA", "CARACTERE_NAO_ASCII_NO_HISTORICO",
            f"{sum(residuos.values())} historico(s) chegam ao destino com caractere "
            "fora da tabela ASCII, porque a lista de 24 letras da planilha nao os cobre. "
            "Confirmar com o cliente se o destino aceita.",
            [f"{c!r} (U+{ord(c):04X}) em {q} lancamentos" for c, q in residuos.most_common()],
        ))
    if por_digitacao:
        ocorrencias.append(Ocorrencia(
            "ALERTA", "RESOLVIDO_POR_DIGITACAO_MANUAL",
            f"{por_digitacao} lancamento(s) so fecharam porque alguem digitou o "
            "codigo a mao numa aba da planilha. Sem CPF, isso nao tem como ser auditado.",
        ))

    ocorrencias.append(Ocorrencia(
        "INFO", "RESOLUCAO_DE_IDENTIDADE",
        f"por nome: {por_nome} | conta de terceiro: {por_terceiro} "
        f"| digitacao manual: {por_digitacao}",
    ))
    return ocorrencias


# ---------------------------------------------------------------------------
# VALIDACAO -- o portao. Nada sobe sem passar aqui.
# ---------------------------------------------------------------------------


def validar(lancamentos):
    ocorrencias = []

    debito = sum((lc.valor for lc in lancamentos if lc.natureza == "DB"), Decimal(0))
    credito = sum((lc.valor for lc in lancamentos if lc.natureza == "CR"), Decimal(0))
    if debito != credito:
        ocorrencias.append(Ocorrencia(
            "BLOQUEIO", "PARTIDA_DOBRADA",
            f"Debito R$ {debito:,.2f} != Credito R$ {credito:,.2f} "
            f"(diferenca R$ {debito - credito:,.2f})",
        ))
    else:
        ocorrencias.append(Ocorrencia(
            "INFO", "PARTIDA_DOBRADA", f"Debito = Credito = R$ {debito:,.2f}"))

    # PERGUNTA EM ABERTO: o arquivo rotulado como competencia 08/2026 traz
    # lancamentos datados de setembro. Se for rescisao paga no mes seguinte,
    # e legitimo e vira ALERTA. Se nao for, e BLOQUEIO. So o cliente responde.
    competencias = Counter((lc.data.year, lc.data.month) for lc in lancamentos)
    if len(competencias) > 1:
        ocorrencias.append(Ocorrencia(
            "ALERTA", "COMPETENCIA_MISTURADA",
            f"O lote contem {len(competencias)} competencias. Confirmar se e esperado.",
            [f"{m:02d}/{a}: {q} lancamentos, R$ "
             f"{sum(lc.valor for lc in lancamentos if (lc.data.year, lc.data.month) == (a, m)):,.2f}"
             for (a, m), q in sorted(competencias.items())],
        ))

    por_loja = defaultdict(lambda: [Decimal(0), Decimal(0)])
    for lc in lancamentos:
        por_loja[lc.loja_destino][0 if lc.natureza == "DB" else 1] += lc.valor
    desbalanceadas = {l: v for l, v in por_loja.items() if v[0] != v[1]}
    if desbalanceadas:
        ocorrencias.append(Ocorrencia(
            "ALERTA", "LOJA_DESBALANCEADA",
            f"{len(desbalanceadas)} de {len(por_loja)} loja(s) nao fecham isoladamente. "
            "Precisa ser confirmado com a Controladoria se e esperado.",
            [f"loja {l}: DB {v[0]:,.2f} / CR {v[1]:,.2f} / dif {v[0] - v[1]:,.2f}"
             for l, v in sorted(desbalanceadas.items())],
        ))

    sem_valor = [lc for lc in lancamentos if lc.valor == 0]
    if sem_valor:
        ocorrencias.append(Ocorrencia(
            "ALERTA", "VALOR_ZERO", f"{len(sem_valor)} lancamento(s) com valor zero."))
    return ocorrencias


# ---------------------------------------------------------------------------
# ADAPTADOR DE DESTINO
# ---------------------------------------------------------------------------


def formatar_valor_como_planilha(valor: Decimal) -> str:
    """
    Reproduz o formato que a planilha produz hoje: separador ponto e zeros
    finais suprimidos (2467,00 -> '2467'; 143,20 -> '143.2').
    Mantido apenas para provar equivalencia com o processo atual.
    """
    texto = f"{valor:.2f}".rstrip("0").rstrip(".")
    return texto or "0"


def formatar_valor_normalizado(valor: Decimal) -> str:
    """Duas casas, sempre. O que deveria ser enviado."""
    return f"{valor:.2f}"


def gerar_lote(lancamentos, formatador, preservar_zeros_esquerda=True):
    linhas = []
    for lc in lancamentos:
        loja = lc.loja_destino or ""
        if not preservar_zeros_esquerda and loja:
            loja = str(int(loja))
        linhas.append([
            loja,
            "0",
            str(lc.conta),
            lc.data.strftime("%d/%m/%Y"),
            formatador(lc.valor),
            lc.historico,
            lc.codigo_final or "",
            lc.nome_colaborador,
            lc.tipo_gerencial or "",
            lc.natureza,
        ])
    return linhas


# ---------------------------------------------------------------------------
# CONFERENCIA CONTRA O PROCESSO ATUAL
# ---------------------------------------------------------------------------


def _chave(linha):
    return tuple(linha)


def ler_saida_da_planilha(caminho: Path):
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    ws = wb["Importação"]
    saida = []
    for l in ws.iter_rows(min_row=2, values_only=True):
        if l[0] is None:
            continue
        data = l[3].strftime("%d/%m/%Y") if isinstance(l[3], datetime) else str(l[3])
        saida.append([
            str(l[0]), "0", str(l[2]), data, str(l[4]),
            str(l[5] or ""), str(l[6] or ""), str(l[7] or ""),
            str(l[8] or ""), str(l[9] or ""),
        ])
    wb.close()
    return saida


def comparar(gerado, referencia):
    a = Counter(_chave(l) for l in gerado)
    b = Counter(_chave(l) for l in referencia)
    return a - b, b - a


# ---------------------------------------------------------------------------


def imprimir(ocorrencias, titulo):
    print(f"\n{titulo}")
    print("-" * len(titulo))
    ordem = {"BLOQUEIO": 0, "ALERTA": 1, "INFO": 2}
    for oc in sorted(ocorrencias, key=lambda o: ordem.get(o.severidade, 9)):
        marca = {"BLOQUEIO": "[X]", "ALERTA": "[!]", "INFO": "[i]"}.get(oc.severidade, "[ ]")
        print(f"{marca} {oc.codigo}: {oc.mensagem}")
        for d in oc.detalhe[:10]:
            print(f"        - {d}")
        if len(oc.detalhe) > 10:
            print(f"        ... e mais {len(oc.detalhe) - 10}")


def main():
    SAIDA.mkdir(exist_ok=True)
    print("=" * 78)
    print("PoC - Motor de integracao Folha -> ERP  |  competencia 08/2026")
    print("=" * 78)

    lancamentos, oc_leitura = ler_folha_dominio(ARQ_FOLHA)
    print(f"\nLidos da origem ............... {len(lancamentos):>6} lancamentos")

    recuperadas = recuperar_pernas_de_credito(lancamentos)
    lancamentos += recuperadas
    print(f"Pernas de credito recuperadas . {len(recuperadas):>6}")
    print(f"Total a processar ............. {len(lancamentos):>6}")

    base, oc_base = ler_base_colaboradores(ARQ_PLANILHA)
    plano = ler_plano_gerencial(ARQ_PLANOGER)
    manuais = ler_codigos_manuais(ARQ_PLANILHA)
    print(f"Base de colaboradores ......... {len(base):>6} nomes distintos")
    print(f"Plano gerencial ............... {len(plano):>6} contas")
    print(f"Codigos digitados a mao ....... {len(manuais):>6}")

    oc_enriq = enriquecer(lancamentos, base, plano, manuais)
    oc_valid = validar(lancamentos)
    todas = oc_leitura + oc_base + oc_enriq + oc_valid
    imprimir(todas, "RELATORIO DE VALIDACAO")

    gerado = gerar_lote(lancamentos, formatar_valor_como_planilha)
    referencia = ler_saida_da_planilha(ARQ_PLANILHA)

    print("\nCONFERENCIA CONTRA A PLANILHA ATUAL")
    print("-" * 35)
    print(f"Linhas geradas pela PoC ....... {len(gerado):>6}")
    print(f"Linhas na aba 'Importacao' .... {len(referencia):>6}")
    so_poc, so_planilha = comparar(gerado, referencia)
    print(f"So na PoC ..................... {sum(so_poc.values()):>6}")
    print(f"So na planilha ................ {sum(so_planilha.values()):>6}")

    if not so_poc and not so_planilha:
        print("\n>>> PARIDADE TOTAL. A PoC reproduz o processo atual linha a linha.")
    else:
        with open(SAIDA / "divergencias.csv", "w", encoding="utf-8-sig", newline="") as fh:
            w = csv.writer(fh, delimiter=";")
            w.writerow(["origem", "loja", "nulo", "conta", "data", "valor",
                        "historico", "codigo", "nome", "tipo", "dc"])
            for k in list(so_poc)[:300]:
                w.writerow(["SO_NA_POC", *k])
            for k in list(so_planilha)[:300]:
                w.writerow(["SO_NA_PLANILHA", *k])
        print(f"\n>>> Divergencias gravadas em {SAIDA / 'divergencias.csv'}")

    with open(SAIDA / "Importar_Folha_API.csv", "w", encoding="utf-8-sig", newline="") as fh:
        csv.writer(fh, delimiter=";").writerows(
            gerar_lote(lancamentos, formatar_valor_normalizado))
    print(f">>> Lote normalizado gravado em {SAIDA / 'Importar_Folha_API.csv'}")

    bloqueios = [o for o in todas if o.severidade == "BLOQUEIO"]
    print("\n" + "=" * 78)
    print(f"VEREDITO: {'CARGA BLOQUEADA' if bloqueios else 'CARGA LIBERADA'} "
          f"({len(bloqueios)} bloqueio(s))")
    print("=" * 78)
    return 1 if bloqueios else 0


if __name__ == "__main__":
    sys.exit(main())
